import json
import streamlit as st
from datetime import datetime
from zoneinfo import ZoneInfo

import pandas as pd
import geopandas as gpd
import ast

import requests
import folium
from streamlit_folium import st_folium


from utils.auth import check_password
from utils.geometry_tools import draw_geometry
from utils.quickscan_tools import save_quickscan
from utils.supabase_client import supabase
from utils.species_groups import SPECIES_GROUPS
from utils.media_tools import save_photos
from utils.gis_analyse import voer_gis_analyse_uit







if check_password():

    st.title("Quickscan",text_alignment="center",)

    tab_create, tab_edit, tab_view = st.tabs(["Nieuwe Quickscan", "Quickscan bewerken", "Quickscan bekijken"])

    with tab_create:
        st.subheader("Nieuwe Quickscan",text_alignment="center",)

        geojson = draw_geometry(key="new_qs_map")

        naam = st.text_input("Projectnaam")
        opmerking = st.text_area("Beschrijving")

        st.subheader("Veldgegevens")
        
        datum = st.date_input("Datum")
        veldwerker = st.text_input("Veldwerker")
        now_local = datetime.now(ZoneInfo("Europe/Amsterdam")).time()
        starttijd = st.time_input("Starttijd",value=now_local)
        eindtijd = st.time_input("Eindtijd",value=now_local)
        
        temperatuur = st.number_input("Temperatuur (°C)", step=0.1)
        windsnelheid = st.selectbox(
            "Windsnelheid",
            ["0 - Stil", "1 - Zwak", "2 - Matig", "3 - Vrij krachtig", "4 - Sterk", "5 - Storm"]
        )
        regen = st.selectbox(
            "Regen",
            ["Geen", "Licht", "Matig", "Hevig"]
        )

        st.subheader("Soortgeschiktheid")
        
        soorten_results = {}
        
        for group, species_list in SPECIES_GROUPS.items():
            suitable = st.toggle(f"Is het gebied geschikt voor {group}?")
        
            if suitable:
                if species_list:
                    selected = st.multiselect(
                        f"Welke soorten binnen {group}?",
                        species_list,
                        key=f"species_{group}"
                    )
                    soorten_results[group] = selected
                else:
                    soorten_results[group] = True
            else:
                soorten_results[group] = False


        st.subheader("Foto's")

        beschikbare_soortgroepen = []
        
        for group, value in soorten_results.items():
            if value is True:
                # Group was marked suitable but has no subspecies
                beschikbare_soortgroepen.append(group)
        
            elif isinstance(value, list) and len(value) > 0:
                # Add each selected species inside the group
                for species in value:
                    beschikbare_soortgroepen.append(species)

                
                
        photos_to_upload = []
        
        add_photo = st.toggle("Wil je een foto toevoegen?", key="add_photo_first")
        
        while add_photo:
            index = len(photos_to_upload)
        
            photo = st.file_uploader(
                "Upload een foto",
                type=["jpg", "jpeg", "png"],
                key=f"photo_{index}"
            )
        
            if photo:
                beschrijving = st.text_area(
                    "Schrijf een beschrijving voor deze foto",
                    key=f"beschrijving_{index}"
                )
        
                soortgroepen_foto = st.multiselect(
                    "Voor welke soortgroepen is dit habitat potentieel geschikt?",
                    beschikbare_soortgroepen,
                    key=f"soortgroepen_{index}"
                )
        
                photos_to_upload.append({
                    "file": photo,
                    "beschrijving": beschrijving,
                    "soortgroepen": soortgroepen_foto
                })
        
                st.success("Foto toegevoegd.")
        
            # IMPORTANT: give this toggle a UNIQUE key
            add_photo = st.toggle(
                "Nog een foto toevoegen?",
                key=f"add_photo_{index}"
            )




        if st.button("Project opslaan"):
        
            if not geojson:
                st.error("Teken eerst een geometrie.")
                st.stop()
        
            if not naam:
                st.error("Voer een projectnaam in.")
                st.stop()
        
            # Check if project name already exists
            existing = supabase.table("new_app_quickscan").select("naam").eq("naam", naam.replace(" ", "_")).execute()
            if existing.data:
                st.error("Projectnaam bestaat al. Kies een andere naam.")
                st.stop()
        
            with st.spinner("Quickscan wordt opgeslagen…"):
            
                qs_id = save_quickscan(
                    naam,
                    opmerking,
                    geojson,
                    datum,
                    veldwerker,
                    starttijd,
                    eindtijd,
                    temperatuur,
                    windsnelheid,
                    regen,
                    soorten_results
                )

            
                if photos_to_upload:
                    save_photos(qs_id, naam, photos_to_upload)
            
            st.success("Quickscan succesvol opgeslagen!")


    with tab_edit:
        st.subheader("Quickscan bewerken", text_alignment="center")
        
        # 1. Haal alle quickscans op
        resp = supabase.table("new_app_quickscan").select("id, naam, datum").execute()
        quickscans = resp.data
        
        if not quickscans:
            st.info("Geen quickscans gevonden.")
            st.stop()
        
        # 2. Dropdown
        keuze = st.selectbox(
            "Kies een Quickscan:",
            options=quickscans,
            format_func=lambda x: f"{x['naam']} – {x['datum']}"
        )

        qs_id = keuze["id"]

        qs_resp = supabase.table("new_app_quickscan") \
            .select("*") \
            .eq("id", qs_id) \
            .single() \
            .execute()
    
        qs = qs_resp.data
    
        if not qs:
            st.error("Quickscan niet gevonden.")
            st.stop()
        


        
        # -----------------------------
        # 1. Geometrie laden
        # -----------------------------
        # Huidige geometrie ophalen
        geometry_data = supabase.storage.from_("new_app").download(qs["geometry_path"])
        st.write(geometry_data)
        current_geojson = geometry_data
        
        # Teken nieuwe geometrie (optioneel)
        st.write("Huidige geometrie:")
        
        st.write("Nieuwe geometrie tekenen (optioneel):")
        nieuwe_geojson = draw_geometry(key="edit_qs_map")
        
        # -----------------------------
        # 2. Basisgegevens
        # -----------------------------
        naam = st.text_input("Projectnaam", value=qs["naam"],key=f"naam_edit_{qs['id']}")
        opmerking = st.text_area("Beschrijving", value=qs["opmerking"],key=f"beschrijving_edit_{qs['id']}")
        
        # -----------------------------
        # 3. Veldgegevens
        # -----------------------------
        st.subheader("Veldgegevens")
        
        datum = st.date_input("Datum", value=qs["datum"],key=f"datum_edit_{qs['id']}")
        veldwerker = st.text_input("Veldwerker", value=qs["veldwerker"],key=f"veldwerker_edit_{qs['id']}")
        
        starttijd = st.time_input("Starttijd", value=qs.get("starttijd"),key=f"starttijd_edit_{qs['id']}")
        eindtijd = st.time_input("Eindtijd", value=qs.get("eindtijd"),key=f"eindtijd_edit_{qs['id']}")
        
        temperatuur = st.number_input("Temperatuur (°C)", step=0.1, value=qs["temperatuur"],key=f"temperatuur_edit_{qs['id']}")
        
        windsnelheid = st.selectbox(
            "Windsnelheid",
            ["0 - Stil", "1 - Zwak", "2 - Matig", "3 - Vrij krachtig", "4 - Sterk", "5 - Storm"],
            index=["0 - Stil", "1 - Zwak", "2 - Matig", "3 - Vrij krachtig", "4 - Sterk", "5 - Storm"].index(qs["windsnelheid"]),
            key=f"windsnelheid_edit_{qs['id']}"
        )
        
        regen = st.selectbox(
            "Regen",
            ["Geen", "Licht", "Matig", "Hevig"],
            index=["Geen", "Licht", "Matig", "Hevig"].index(qs["regen"]),
            key=f"regen_{qs['id']}"
        )
        
        # -----------------------------
        # 4. Soortgeschiktheid
        # -----------------------------
        st.subheader("Soortgeschiktheid")
        
        soorten_results = qs["soorten"]  # dit is jouw JSON dict uit de database
        nieuwe_soorten = {}
        
        for group, species_list in SPECIES_GROUPS.items():
            huidige_waarde = soorten_results.get(group)
        
            suitable = st.toggle(
                f"Is het gebied geschikt voor {group}?",
                value=(huidige_waarde is True or isinstance(huidige_waarde, list)),
                key=f"suitable_{group}_{qs['id']}"
            )
        
            if suitable:
                if species_list:
                    selected = st.multiselect(
                        f"Welke soorten binnen {group}?",
                        species_list,
                        default=huidige_waarde if isinstance(huidige_waarde, list) else [],
                        key=f"species_{group}_{qs['id']}"
                    )
                    nieuwe_soorten[group] = selected
                else:
                    nieuwe_soorten[group] = True
            else:
                nieuwe_soorten[group] = False
        
        # -----------------------------
        # 5. Foto's beheren
        # -----------------------------
        st.subheader("Foto's")
        
        fotos_resp = supabase.table("new_app_quickscan_fotos") \
            .select("*") \
            .eq("quickscan_id", qs["id"]) \
            .execute()
        
        fotos = fotos_resp.data
        
        # Build available species/groups based on Quickscan suitability
        soorten_results = nieuwe_soorten
        beschikbare_soortgroepen = []
        
        for group, value in soorten_results.items():
            if value is True:
                beschikbare_soortgroepen.append(group)
            elif isinstance(value, list) and len(value) > 0:
                beschikbare_soortgroepen.extend(value)
        
        for foto in fotos:
            import os
            filename = os.path.basename(foto["foto_pad"])
        
        
            data = supabase.storage.from_("new_app").download(foto["foto_pad"])
            st.image(data)
    
            nieuwe_beschrijving = st.text_area(
                f"Beschrijving",
                value=foto.get("beschrijving", ""),
                key=f"beschrijving_{qs['id']}_{foto['id']}"
            )
    
            nieuwe_soorten_list = st.multiselect(
                f"Soortgroepen",
                beschikbare_soortgroepen,
                key=f"soorten_{qs['id']}_{foto['id']}"
            )
            
            col1, col2 = st.columns([1, 1])
            if col1.button("Opslaan wijzigingen", key=f"save_{qs['id']}_{foto['id']}"):
                supabase.table("new_app_quickscan_fotos") \
                    .update({
                        "beschrijving": nieuwe_beschrijving,
                        "soortgroep": nieuwe_soorten_list
                    }) \
                    .eq("id", foto["id"]) \
                    .execute()
    
                st.success("Foto metadata bijgewerkt.")
                st.rerun()
    
            if col2.button(f"Verwijder", key=f"delete_{qs['id']}_{foto['id']}"):
                supabase.storage.from_("new_app").remove(foto["foto_pad"])
                supabase.table("new_app_quickscan_fotos") \
                    .delete() \
                    .eq("id", foto["id"]) \
                    .execute()
    
                st.warning(f"verwijderd.")
                st.rerun()

            "---"


        if "upload_key" not in st.session_state:
            st.session_state.upload_key = f"upload_{qs['id']}"
        
        nieuwe_foto = st.file_uploader(
            "Nieuwe foto uploaden",
            type=["jpg", "jpeg", "png"],
            key=st.session_state.upload_key
        )

        
        if nieuwe_foto:
            import uuid
            unique_id = str(uuid.uuid4())
            safe_name = naam.replace(" ", "_")
        
            filename = f"quickscan/fotos/{unique_id}.jpg"
        
            # Upload to storage
            supabase.storage.from_("new_app").upload(
                filename,
                nieuwe_foto.read(),
                file_options={"content-type": "image/jpeg", "x-upsert": "true"}
            )
        
            # Insert metadata row
            supabase.table("new_app_quickscan_fotos").insert({
                "quickscan_id": qs["id"],
                "quickscan_naam": safe_name,
                "foto_pad": filename,
                "beschrijving": "",
                "soortgroep": []
            }).execute()
        
            st.success("Foto geüpload.")
            
           # Reset uploader key to clear the file
            st.session_state.upload_key = f"upload_{qs['id']}_{uuid.uuid4()}"
            
            st.rerun()

        
        # -----------------------------
        # 6. Opslaan
        # -----------------------------
        if st.button("Quickscan opslaan"):
            update_data = {
                "naam": naam,
                "opmerking": opmerking,
                "datum": str(datum),
                "veldwerker": veldwerker,
                "starttijd": str(starttijd),
                "eindtijd": str(eindtijd),
                "temperatuur": temperatuur,
                "windsnelheid": windsnelheid,
                "regen": regen,
                "soorten": nieuwe_soorten,
            }
        
            # Geometrie vervangen indien nieuwe getekend
            if nieuwe_geojson:
                path = f"geometry/{qs['id']}.geojson"
                supabase.storage.from_("new_app").upload(path, nieuwe_geojson)
                update_data["geometry_path"] = path
        
            supabase.table("new_app_quickscan").update(update_data).eq("id", qs["id"]).execute()
        
            st.success("Quickscan bijgewerkt.")
            st.session_state.edit_mode = False
            st.rerun()





    

    
    with tab_view:
        st.header("Quickscan bekijken",text_alignment="center",)
    
        # Load all quickscans
        all_qs = supabase.table("new_app_quickscan") \
            .select("naam, datum") \
            .order("datum", desc=True) \
            .execute()
    
        if not all_qs.data:
            st.info("Nog geen Quickscans beschikbaar.")
            st.stop()
    
        # Dropdown
        qs_labels = [f"{qs['naam']} ({qs['datum']})" for qs in all_qs.data]
        selected_label = st.selectbox("Kies een Quickscan:", qs_labels)
    
        selected_name = selected_label.split(" (")[0]
    
        # Load selected Quickscan
        qs = supabase.table("new_app_quickscan") \
            .select("*") \
            .eq("naam", selected_name) \
            .single() \
            .execute().data
    
        # Show info
        st.subheader("Projectinformatie")
        st.write(f"**Naam:** {qs['naam']}")
        st.write(f"**Datum:** {qs['datum']}")
        st.write(f"**Tijd:** {qs['starttijd']} - {qs['eindtijd']}")
        st.write(f"**Veldwerker:** {qs['veldwerker']}")
        st.write(f"**Beschrijving:** {qs['opmerking']}")
    
        st.subheader("Weersomstandigheden")
        st.write(f"**Temperatuur:** {qs['temperatuur']} °C")
        st.write(f"**Wind:** {qs['windsnelheid']} Bft")
        st.write(f"**Regen:** {qs['regen']}")
        
        "---"
        st.subheader("Soortgeschiktheid")

        
        soorten = qs["soorten"]
        
        rows = []
        
        for group, value in soorten.items():
            if value is True:
                rows.append([group, "Geschikt", "—"])
            elif value is False:
                rows.append([group, "Niet geschikt", "—"])
            elif isinstance(value, list):
                rows.append([group, "Geschikt", ", ".join(value)])
        
        df = pd.DataFrame(rows, columns=["Soortgroep", "Geschiktheid", "Soorten"])
        
        st.table(df)
        
    
        "---"
        st.subheader("Foto's")
        fotos = supabase.table("new_app_quickscan_fotos") \
            .select("*") \
            .eq("quickscan_id", qs["id"]) \
            .execute()
        
        if not fotos.data:
            st.info("Geen foto's toegevoegd.")
        else:
            import ast
        
            for foto in fotos.data:
                url = supabase.storage.from_("new_app").get_public_url(foto["foto_pad"])
                st.image(url, caption=foto["beschrijving"])
        
                soorten_list = foto["soortgroep"]
        
                # Convert JSON string → Python list
                if isinstance(soorten_list, str):
                    soorten_list = ast.literal_eval(soorten_list)
        
                st.write(f"**Potentieel geschikt voor:** {', '.join(soorten_list)}")
               

        "---"
        st.markdown("""
        ### Extra GIS‑analyse: afstand tot Natura 2000‑gebieden
        
        Met deze aanvullende analyse wordt gecontroleerd of het onderzoeksgebied zich binnen een straal van **3 kilometer** van een Natura 2000‑gebied bevindt.  
        De tool berekent automatisch:
        
        - het **centroid** van het Quickscan‑gebied  
        - een **buffer van 3 km** rond dit punt  
        - de **intersectie** met alle Natura 2000‑gebieden  
        - een **kaartvisualisatie** met alle relevante lagen  
        
        Klik op de knop hieronder om de analyse uit te voeren.
        """)

        run_analysis = st.toggle(
            "Wil je controleren of het onderzoeksgebied binnen 3 km van een Natura 2000‑gebied ligt?"
        )


        if run_analysis:
            voer_gis_analyse_uit(qs, supabase)

        "---"
        st.markdown("### Conclusie van de Quickscan")
        
        # 1. Vraag of gebruiker een conclusie wil schrijven
        write_conclusion = st.toggle("Wil je een conclusie toevoegen of bewerken?")
        
        # Haal bestaande conclusie op (indien aanwezig)
        existing = supabase.table("new_app_quickscan_conclusions") \
            .select("*") \
            .eq("project_naam", qs["naam"]) \
            .execute()
        
        existing_text = None
        existing_id = None
        
        if existing.data:
            existing_text = existing.data[0]["conclusie"]
            existing_id = existing.data[0]["id"]

        if write_conclusion:
                
            st.markdown("#### Schrijf of bewerk de conclusie")
        
            # Tekstvak met bestaande tekst indien aanwezig
            conclusion_text = st.text_area(
                "Conclusie:",
                value=existing_text if existing_text else "",
                placeholder="..."
            )
        
            # Als er al een conclusie bestaat → toon update-knop
            if existing_id:
        
                st.info("Er is al een conclusie opgeslagen voor dit project.")
        
                col1, col2, col3 = st.columns(3)
        
                # --- Bijwerken ---
                with col1:
                    if st.button("Conclusie bijwerken"):
                        if conclusion_text.strip() == "":
                            st.error("De conclusie mag niet leeg zijn.")
                        else:
                            supabase.table("new_app_quickscan_conclusions") \
                                .update({"conclusie": conclusion_text}) \
                                .eq("id", existing_id) \
                                .execute()
        
                            st.success("De conclusie is bijgewerkt.")
                            st.rerun()   # 🔄 reload
        
                # --- Verwijderen ---
                with col2:
                    if st.button("Conclusie verwijderen"):
                        supabase.table("new_app_quickscan_conclusions") \
                            .delete() \
                            .eq("id", existing_id) \
                            .execute()
        
                        st.warning("De conclusie is verwijderd.")
                        st.rerun()   # 🔄 reload
        
                # --- Handmatige reload ---
                with col3:
                    if st.button("Vernieuwen"):
                        st.rerun()   # 🔄 reload
        
            # Als er nog GEEN conclusie bestaat → toon opslaan-knop
            else:
                if st.button("Nieuwe conclusie opslaan"):
                    if conclusion_text.strip() == "":
                        st.error("De conclusie mag niet leeg zijn.")
                    else:
                        supabase.table("new_app_quickscan_conclusions") \
                            .insert({
                                "project_naam": qs["naam"],
                                "conclusie": conclusion_text
                            }).execute()
        
                        st.success("De conclusie is succesvol opgeslagen.")
                        st.rerun()   # 🔄 reload

 
        "---"
        from docx import Document
        from docx.shared import Inches
        import io
        
        st.subheader("Download Quickscan Rapport (.docx)")
        
        # Haal conclusie op
        conclusion_data = supabase.table("new_app_quickscan_conclusions") \
            .select("*") \
            .eq("project_naam", qs["naam"]) \
            .execute()
        
        conclusion_text = (
            conclusion_data.data[0]["conclusie"]
            if conclusion_data.data else "Geen conclusie opgeslagen."
        )
        
        # Maak Word-document
        doc = Document()
        
        doc.add_heading("Quickscan Rapport", level=1)
        
        doc.add_heading("Projectinformatie", level=2)
        doc.add_paragraph(f"Naam: {qs['naam']}")
        doc.add_paragraph(f"Datum: {qs['datum']}")
        doc.add_paragraph(f"Veldwerker: {qs['veldwerker']}")
        doc.add_paragraph(f"Opmerking: {qs['opmerking']}")
        
        doc.add_heading("Weersomstandigheden", level=2)
        doc.add_paragraph(f"Temperatuur: {qs['temperatuur']} °C")
        doc.add_paragraph(f"Wind: {qs['windsnelheid']} Bft")
        doc.add_paragraph(f"Regen: {qs['regen']}")
        

        
        # Voeg foto toe (als aanwezig)
        if "foto_path" in qs:
            img_bytes = supabase.storage.from_("new_app").download(qs["foto_path"])
            doc.add_heading("Foto", level=2)
            doc.add_picture(io.BytesIO(img_bytes), width=Inches(4))
        
        # Document in geheugen opslaan
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        # Download-knop
        st.download_button(
            "Download Quickscan Rapport (.docx)",
            buffer,
            file_name=f"{qs['naam']}_quickscan.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )








